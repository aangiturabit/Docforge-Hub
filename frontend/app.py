import streamlit as st
import requests
import json
import io
from datetime import datetime, date

API_BASE = "http://127.0.0.1:8000/api"

st.set_page_config(
    page_title="DocForge Hub",
    page_icon="📄",
    layout="wide",
    initial_sidebar_state="expanded"
)


# ═══════════════════════════════════════════
# HELPERS
# ═══════════════════════════════════════════

def short_name(name: str, max_len: int = 38) -> str:
    if len(name) <= max_len:
        return name
    truncated = name[:max_len]
    last_space = truncated.rfind(" ")
    if last_space > 20:
        truncated = truncated[:last_space]
    return truncated + "…"


def api_get(path, params=None, show_error=True):
    try:
        r = requests.get(f"{API_BASE}{path}", params=params, timeout=30)
        if r.status_code == 200:
            return r.json()
        if show_error:
            st.error(f"API error {r.status_code} on GET {path}: {r.text[:300]}")
        return None
    except requests.exceptions.ConnectionError:
        if show_error:
            st.error(f"Cannot connect to API at {API_BASE}. Is the backend running?")
        return None
    except Exception as e:
        if show_error:
            st.error(f"Unexpected error on GET {path}: {str(e)}")
        return None


def api_post(path, body, show_error=True):
    try:
        r = requests.post(f"{API_BASE}{path}", json=body, timeout=120)
        if r.status_code == 200:
            return r.json(), None
        err_msg = r.text[:300]
        if show_error:
            st.error(f"API error {r.status_code} on POST {path}: {err_msg}")
        return None, err_msg
    except requests.exceptions.ConnectionError:
        msg = f"Cannot connect to API at {API_BASE}"
        return None, msg
    except Exception as e:
        return None, str(e)


def api_delete(path, show_error=True):
    try:
        r = requests.delete(f"{API_BASE}{path}", timeout=30)
        if r.status_code == 200:
            return True
        if show_error:
            st.error(f"API error {r.status_code} on DELETE {path}: {r.text[:200]}")
        return False
    except Exception as e:
        if show_error:
            st.error(str(e))
        return False


def format_date(dt_str):
    try:
        return datetime.fromisoformat(dt_str).strftime("%d %b %Y")
    except Exception:
        return dt_str or ""


def status_badge(status: str) -> str:
    return {
        "validated":    "🟢",
        "pending":      "🟡",
        "needs_review": "🟠",
        "failed":       "🔴",
        "draft":        "🔵"
    }.get(status, "⚪")


def get_structured_sections(doc_data: dict) -> list:
    if not doc_data:
        return []
    structured_json = doc_data.get("structured_json")
    if structured_json:
        try:
            parsed = json.loads(structured_json) if isinstance(structured_json, str) else structured_json
            if isinstance(parsed, list):
                return parsed
            if isinstance(parsed, dict):
                return parsed.get("sections", [])
        except Exception:
            pass
    return []


def render_document(content: str):
    if not content:
        st.info("No content to display.")
        return
    lines = content.split("\n")
    i = 0
    while i < len(lines):
        line   = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            if table_lines:
                st.markdown("\n".join(table_lines))
            continue
        if "   |   " in stripped:
            parts = [p.strip() for p in stripped.split("|")]
            cols  = st.columns(len(parts))
            for col, part in zip(cols, parts):
                col.write(part)
            i += 1
            continue
        if stripped.startswith("- "):
            st.markdown(stripped)
            i += 1
            continue
        if len(stripped) > 2 and stripped[0].isdigit() and stripped[1] in ".)":
            st.markdown(stripped)
            i += 1
            continue
        if (
            len(stripped) < 65
            and not stripped.startswith("-")
            and not stripped[0].isdigit()
            and len(stripped.split()) <= 9
            and stripped == stripped.rstrip(":")
        ):
            st.markdown(f"**{stripped}**")
            i += 1
            continue
        st.write(stripped)
        i += 1


# ─────────────────────────────────────────
# PDF GENERATION
# ─────────────────────────────────────────

def generate_pdf_from_structured(
    structured_sections: list,
    title: str,
    company: dict = None,
    department: str = ""
) -> bytes:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.lib import colors
        from reportlab.platypus import (
            SimpleDocTemplate, Paragraph, Spacer,
            Table, TableStyle, PageBreak, HRFlowable, KeepTogether
        )
        from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_JUSTIFY
        import re

        def clean(text):
            if not text:
                return ""
            text = re.sub(r"#{1,6}\s*", "", str(text))
            text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
            text = re.sub(r"\*(.*?)\*",     r"\1", text)
            text = text.replace("\u2019", "'").replace("\u2018", "'")
            text = text.replace("\u201c", '"').replace("\u201d", '"')
            text = text.replace("\u2013", "-").replace("\u2014", "--")
            text = text.encode("latin-1", "replace").decode("latin-1")
            return text.strip()

        buffer = io.BytesIO()
        W = A4[0] - 5 * cm
        doc = SimpleDocTemplate(
            buffer, pagesize=A4,
            rightMargin=2.5*cm, leftMargin=2.5*cm,
            topMargin=2.5*cm,   bottomMargin=2.5*cm,
            title=title
        )

        cover_title_style = ParagraphStyle(
            "CT", fontSize=22, fontName="Helvetica-Bold",
            textColor=colors.HexColor("#1a1a1a"),
            alignment=TA_CENTER, spaceAfter=10, leading=28
        )
        cover_sub_style = ParagraphStyle(
            "CS", fontSize=12, fontName="Helvetica",
            textColor=colors.HexColor("#555555"),
            alignment=TA_CENTER, spaceAfter=6
        )
        cover_meta_style = ParagraphStyle(
            "CM", fontSize=10, fontName="Helvetica",
            textColor=colors.HexColor("#888888"),
            alignment=TA_CENTER, spaceAfter=5
        )
        heading_style = ParagraphStyle(
            "SH", fontSize=11, fontName="Helvetica-Bold",
            textColor=colors.HexColor("#1a1a1a"),
            spaceBefore=18, spaceAfter=6,
            backColor=colors.HexColor("#f2f2f2"),
            borderPadding=(5, 8, 5, 8)
        )
        body_style = ParagraphStyle(
            "B", fontSize=10, fontName="Helvetica",
            leading=16, spaceAfter=7, alignment=TA_JUSTIFY
        )
        list_style = ParagraphStyle(
            "L", fontSize=10, fontName="Helvetica",
            leading=14, spaceAfter=4, leftIndent=18, firstLineIndent=-10
        )

        story = []

        # Cover
        story.append(Spacer(1, 3 * cm))
        company_name = company.get("name", "") if company else ""
        if company_name:
            story.append(Paragraph(clean(company_name), cover_sub_style))
            story.append(Spacer(1, 0.4 * cm))
        story.append(HRFlowable(width=W, thickness=2,
                                 color=colors.HexColor("#1a1a1a"), spaceAfter=14))
        story.append(Paragraph(clean(title), cover_title_style))
        story.append(HRFlowable(width=W, thickness=1,
                                 color=colors.HexColor("#cccccc"),
                                 spaceBefore=14, spaceAfter=18))
        if department:
            story.append(Paragraph(f"Department: {clean(department)}", cover_meta_style))
        if company:
            if company.get("industry"):
                story.append(Paragraph(f"Industry: {clean(company.get('industry',''))}", cover_meta_style))
            if company.get("location"):
                story.append(Paragraph(f"Location: {clean(company.get('location',''))}", cover_meta_style))
        story.append(Spacer(1, 0.4 * cm))
        story.append(Paragraph(f"Generated: {date.today().strftime('%d %B %Y')}", cover_meta_style))
        story.append(Spacer(1, 3.5 * cm))
        story.append(PageBreak())

        align_map = {"left": TA_LEFT, "center": TA_CENTER, "justify": TA_JUSTIFY}

        for section in structured_sections:
            heading      = clean(section.get("heading", ""))
            content_type = section.get("content_type", "text")
            content      = section.get("content", "")
            styling      = section.get("styling", {})
            align        = align_map.get(styling.get("alignment", "justify"), TA_JUSTIFY)

            sec_story = []

            if heading:
                sec_story.append(Paragraph(heading, heading_style))
                sec_story.append(Spacer(1, 0.15 * cm))

            if content_type == "text" and isinstance(content, str):
                para_style = ParagraphStyle("DB", parent=body_style, alignment=align)
                normalized = re.sub(r'(?<!\n)\n(?!\n)', ' ', content)
                for para in normalized.split("\n\n"):
                    c = clean(para.strip())
                    if c and len(c) > 3:
                        sec_story.append(Paragraph(c, para_style))
                        sec_story.append(Spacer(1, 0.15 * cm))

            elif content_type == "table" and isinstance(content, list) and content:
                table_data = [
                    [clean(str(c)) for c in row.get("cells", [])]
                    for row in content
                ]
                if table_data and table_data[0]:
                    col_count = len(table_data[0])
                    col_width = W / col_count
                    t = Table(table_data, colWidths=[col_width] * col_count, repeatRows=1)
                    t.setStyle(TableStyle([
                        ("BACKGROUND",    (0, 0), (-1, 0),  colors.HexColor("#2c2c2c")),
                        ("TEXTCOLOR",     (0, 0), (-1, 0),  colors.white),
                        ("FONTNAME",      (0, 0), (-1, 0),  "Helvetica-Bold"),
                        ("FONTSIZE",      (0, 0), (-1, -1), 9),
                        ("TOPPADDING",    (0, 0), (-1, -1), 5),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                        ("LEFTPADDING",   (0, 0), (-1, -1), 7),
                        ("RIGHTPADDING",  (0, 0), (-1, -1), 7),
                        ("FONTNAME",      (0, 1), (-1, -1), "Helvetica"),
                        ("ROWBACKGROUNDS",(0, 1), (-1, -1), [colors.white, colors.HexColor("#f7f7f7")]),
                        ("GRID",          (0, 0), (-1, -1), 0.4, colors.HexColor("#dddddd")),
                        ("ALIGN",         (0, 0), (-1, -1), "LEFT"),
                        ("VALIGN",        (0, 0), (-1, -1), "MIDDLE"),
                    ]))
                    sec_story.append(t)
                    sec_story.append(Spacer(1, 0.3 * cm))

            elif content_type == "list" and isinstance(content, list):
                for item in content:
                    c = clean(str(item))
                    if c:
                        sec_story.append(Paragraph(f"\u2022  {c}", list_style))
                sec_story.append(Spacer(1, 0.2 * cm))

            if styling.get("page_break_after"):
                sec_story.append(PageBreak())
            else:
                sec_story.append(Spacer(1, 0.2 * cm))

            if sec_story:
                story.append(KeepTogether(sec_story[:3]))
                story.extend(sec_story[3:])

        doc.build(story)
        buffer.seek(0)
        return buffer.read()
    except Exception as e:
        st.error(f"PDF generation error: {str(e)}")
        return None


# ─────────────────────────────────────────
# DOCX GENERATION
# ─────────────────────────────────────────

def generate_docx_from_structured(
    structured_sections: list,
    title: str,
    company: dict = None,
    department: str = ""
) -> bytes:
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        import re

        def clean(text):
            if not text:
                return ""
            text = re.sub(r"#{1,6}\s*", "", str(text))
            text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
            text = re.sub(r"\*(.*?)\*",     r"\1", text)
            text = text.replace("\u2019", "'").replace("\u2018", "'")
            text = text.replace("\u201c", '"').replace("\u201d", '"')
            text = text.replace("\u2013", "-").replace("\u2014", "--")
            return text.strip()

        doc = Document()
        for sec in doc.sections:
            sec.top_margin    = Cm(2.5)
            sec.bottom_margin = Cm(2.5)
            sec.left_margin   = Cm(2.5)
            sec.right_margin  = Cm(2.5)

        def add_para(text, bold=False, size=10,
                     align=WD_ALIGN_PARAGRAPH.LEFT,
                     sb=0, sa=6, color=None):
            p = doc.add_paragraph()
            p.alignment = align
            p.paragraph_format.space_before = Pt(sb)
            p.paragraph_format.space_after  = Pt(sa)
            run = p.add_run(clean(text))
            run.bold      = bold
            run.font.size = Pt(size)
            if color:
                run.font.color.rgb = RGBColor(*color)
            return p

        def add_hr():
            p    = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(10)
            pPr  = p._p.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            bot  = OxmlElement("w:bottom")
            bot.set(qn("w:val"),   "single")
            bot.set(qn("w:sz"),    "12")
            bot.set(qn("w:space"), "1")
            bot.set(qn("w:color"), "1a1a1a")
            pBdr.append(bot)
            pPr.append(pBdr)

        # Cover
        doc.add_paragraph()
        doc.add_paragraph()
        company_name = company.get("name", "") if company else ""
        if company_name:
            add_para(company_name, size=13, align=WD_ALIGN_PARAGRAPH.CENTER,
                     sa=6, color=(68, 68, 68))
        add_hr()
        add_para(title, bold=True, size=20,
                 align=WD_ALIGN_PARAGRAPH.CENTER,
                 sb=10, sa=10, color=(26, 26, 26))
        add_hr()
        if department:
            add_para(f"Department: {department}", size=10,
                     align=WD_ALIGN_PARAGRAPH.CENTER, sa=3, color=(100, 100, 100))
        if company and company.get("industry"):
            add_para(f"Industry: {company.get('industry', '')}", size=10,
                     align=WD_ALIGN_PARAGRAPH.CENTER, sa=3, color=(100, 100, 100))
        add_para(f"Generated: {date.today().strftime('%d %B %Y')}", size=10,
                 align=WD_ALIGN_PARAGRAPH.CENTER, sa=3, color=(136, 136, 136))
        doc.add_page_break()

        align_map = {
            "left":    WD_ALIGN_PARAGRAPH.LEFT,
            "center":  WD_ALIGN_PARAGRAPH.CENTER,
            "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        }

        for section_data in structured_sections:
            heading      = clean(section_data.get("heading", ""))
            content_type = section_data.get("content_type", "text")
            content      = section_data.get("content", "")
            styling      = section_data.get("styling", {})
            para_align   = align_map.get(
                styling.get("alignment", "justify"),
                WD_ALIGN_PARAGRAPH.JUSTIFY
            )

            if heading:
                h = doc.add_paragraph()
                h.paragraph_format.space_before = Pt(16)
                h.paragraph_format.space_after  = Pt(6)
                pPr = h._p.get_or_add_pPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"),   "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"),  "F2F2F2")
                pPr.append(shd)
                run = h.add_run(heading)
                run.bold           = True
                run.font.size      = Pt(11)
                run.font.color.rgb = RGBColor(26, 26, 26)

            if content_type == "text" and isinstance(content, str):
                normalized = re.sub(r'(?<!\n)\n(?!\n)', ' ', content)
                for para_text in normalized.split("\n\n"):
                    c = clean(para_text.strip())
                    if not c:
                        continue
                    p = doc.add_paragraph()
                    p.alignment                     = para_align
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after  = Pt(6)
                    p.paragraph_format.line_spacing = Pt(15)
                    run = p.add_run(c)
                    run.font.size = Pt(10)

            elif content_type == "table" and isinstance(content, list) and content:
                rows = len(content)
                cols = max((len(r.get("cells", [])) for r in content), default=0)
                if rows > 0 and cols > 0:
                    table = doc.add_table(rows=rows, cols=cols)
                    table.style = "Table Grid"
                    for i, row_data in enumerate(content):
                        cells = row_data.get("cells", [])
                        for j in range(cols):
                            cell_text = cells[j] if j < len(cells) else ""
                            cell      = table.cell(i, j)
                            cell.text = clean(str(cell_text))
                            p         = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(3)
                            p.paragraph_format.space_after  = Pt(3)
                            if p.runs:
                                p.runs[0].font.size = Pt(9)
                                if i == 0:
                                    p.runs[0].bold = True
                                    tc_pr = cell._tc.get_or_add_tcPr()
                                    shd   = OxmlElement("w:shd")
                                    shd.set(qn("w:val"),   "clear")
                                    shd.set(qn("w:color"), "auto")
                                    shd.set(qn("w:fill"),  "2C2C2C")
                                    tc_pr.append(shd)
                                    p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
                    doc.add_paragraph().paragraph_format.space_after = Pt(8)

            elif content_type == "list" and isinstance(content, list):
                for item in content:
                    c = clean(str(item))
                    if c:
                        p = doc.add_paragraph(style="List Bullet")
                        p.paragraph_format.space_before = Pt(1)
                        p.paragraph_format.space_after  = Pt(3)
                        run = p.add_run(c)
                        run.font.size = Pt(10)

            doc.add_paragraph().paragraph_format.space_after = Pt(4)

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.read()
    except Exception as e:
        st.error(f"DOCX generation error: {str(e)}")
        return None


# ═══════════════════════════════════════════
# SESSION STATE
# ═══════════════════════════════════════════

_defaults = {
    "page":            "generate",
    "session_id":      None,
    "sections":        None,
    "department_id":   None,
    "template_id":     None,
    "template_name":   None,
    "department_name": None,
    "document":        None,
    "preview_content": None,
    "answers":         {},
    "show_regen":      False,
    "library_doc":     None,
    "val_data":        None,   # stores last validation result
    "company": {
        "name": "", "industry": "",
        "size": "", "location": "", "tone": "Professional"
    }
}

for k, v in _defaults.items():
    if k not in st.session_state:
        st.session_state[k] = v


# ═══════════════════════════════════════════
# SIDEBAR
# ═══════════════════════════════════════════

with st.sidebar:
    st.markdown("## 📄 DocForge Hub")
    st.caption("AI-Powered Document Generator")
    st.divider()

    dept_data  = api_get("/departments", show_error=False) or []
    dept_names = [d["name"] for d in dept_data]
    dept_map   = {d["name"]: d["id"] for d in dept_data}
    selected_dept = st.selectbox("Department", dept_names, key="sb_dept")

    selected_template = None
    tmpl_map          = {}
    tmpl_display_map  = {}
    if selected_dept:
        dept_id   = dept_map.get(selected_dept)
        tmpl_data = api_get(f"/templates/{dept_id}", show_error=False) or []
        tmpl_display_map   = {short_name(t["name"]): t["name"] for t in tmpl_data}
        tmpl_map           = {t["name"]: t["id"] for t in tmpl_data}
        tmpl_display_names = [short_name(t["name"]) for t in tmpl_data]
        selected_tmpl_display = st.selectbox("Document Type", tmpl_display_names, key="sb_tmpl")
        selected_template = tmpl_display_map.get(selected_tmpl_display) if selected_tmpl_display else None

    st.divider()
    st.markdown("**Navigation**")
    if st.session_state["page"] == "generate":
        if st.button("📚 Document Library", use_container_width=True):
            st.session_state["page"]        = "library"
            st.session_state["library_doc"] = None
            st.rerun()
    else:
        if st.button("← Back to Generator", use_container_width=True):
            st.session_state["page"] = "generate"
            st.rerun()

    st.divider()
    with st.expander("🏢 Company Context", expanded=False):
        st.caption("Applied to all documents")
        st.session_state["company"]["name"]     = st.text_input("Company Name",   value=st.session_state["company"]["name"],     key="co_name")
        st.session_state["company"]["industry"] = st.text_input("Industry",       value=st.session_state["company"]["industry"], key="co_industry")
        sizes = ["", "1-10", "11-50", "51-200", "201-500", "500+"]
        st.session_state["company"]["size"]     = st.selectbox(
            "Company Size", sizes,
            index=sizes.index(st.session_state["company"].get("size", ""))
            if st.session_state["company"].get("size", "") in sizes else 0,
            key="co_size"
        )
        st.session_state["company"]["location"] = st.text_input("Location",       value=st.session_state["company"]["location"], key="co_location")
        tones = ["Professional", "Formal", "Friendly", "Technical", "Empathetic"]
        st.session_state["company"]["tone"]     = st.selectbox(
            "Document Tone", tones,
            index=tones.index(st.session_state["company"].get("tone", "Professional")),
            key="co_tone"
        )

    st.divider()
    st.caption("DocForge Hub v1.0")


# ═══════════════════════════════════════════
# PAGE: GENERATE
# ═══════════════════════════════════════════

if st.session_state["page"] == "generate":

    doc     = st.session_state.get("document")
    preview = st.session_state.get("preview_content")

    # ── No state yet — welcome ───────────────────────────────────────────────
    if not st.session_state.get("sections") and not doc and not preview:
        st.markdown("### 📄 DocForge Hub")
        st.caption("Select a department and document type from the sidebar, then click Generate Questions.")
        st.write("")
        if st.button("🚀 Generate Questions", type="primary", use_container_width=False, key="gen_q_init"):
            if selected_template:
                with st.spinner("Generating questions…"):
                    data, err = api_post(
                        "/generate/questions",
                        {"document_type_id": tmpl_map.get(selected_template)}
                    )
                    if data:
                        st.session_state.update({
                            "session_id":    data["session_id"],
                            "sections":      data["sections"],
                            "department_id": dept_map[selected_dept],
                            "template_id":   tmpl_map[selected_template],
                            "template_name": selected_template,
                            "department_name": selected_dept,
                            "document":      None,
                            "preview_content": None,
                            "answers":       {},
                            "show_regen":    False,
                            "val_data":      None,
                        })
                        st.rerun()
                    else:
                        st.error(f"Error: {err}")
            else:
                st.warning("Please select a Department and Document Type first.")
        st.stop()

    # ── Two-column layout: LEFT = form/questions, RIGHT = document ───────────
    left_col, right_col = st.columns([2, 3], gap="large")

    # ════════════════════════════════════════
    # LEFT COLUMN — questions + controls
    # ════════════════════════════════════════
    with left_col:

        # Header
        st.markdown(f"### {st.session_state.get('template_name', 'Document')}")
        st.caption(st.session_state.get("department_name", ""))
        st.write("")

        # Generate Questions button (re-trigger)
        if st.button("🚀 Generate Questions", type="secondary",
                     use_container_width=True, key="gen_q_btn"):
            if selected_template:
                with st.spinner("Generating questions…"):
                    data, err = api_post(
                        "/generate/questions",
                        {"document_type_id": tmpl_map.get(selected_template)}
                    )
                    if data:
                        st.session_state.update({
                            "session_id":      data["session_id"],
                            "sections":        data["sections"],
                            "department_id":   dept_map[selected_dept],
                            "template_id":     tmpl_map[selected_template],
                            "template_name":   selected_template,
                            "department_name": selected_dept,
                            "document":        None,
                            "preview_content": None,
                            "answers":         {},
                            "show_regen":      False,
                            "val_data":        None,
                        })
                        st.rerun()
            else:
                st.warning("Select a document type first.")

        # Questions form
        if st.session_state.get("sections"):
            st.divider()
            answers = {}

            for section in st.session_state["sections"]:
                fields = section.get("fields", [])
                if not fields:
                    continue
                st.markdown(f"**{section['section_name']}**")
                for field in fields:
                    fn      = field["field_name"]
                    label   = field["field_label"]
                    ftype   = field["field_type"]
                    req     = field["is_required"]
                    display = f"{label}{'  ✱' if req else ''}"

                    if ftype == "textarea":
                        answers[fn] = st.text_area(display, key=f"f_{fn}", height=90)
                    elif ftype == "date":
                        answers[fn] = str(st.date_input(display, key=f"f_{fn}"))
                    elif ftype == "number":
                        answers[fn] = str(st.number_input(display, key=f"f_{fn}", step=1))
                    else:
                        answers[fn] = st.text_input(display, key=f"f_{fn}")

                st.write("")

            st.session_state["answers"] = answers
            st.divider()

            # Generate Document button
            if st.button("📄 Generate Document", type="primary",
                         use_container_width=True, key="gen_doc_btn"):
                with st.spinner("Generating document…"):
                    data, err = api_post(
                        "/generate/document",
                        {
                            "session_id":    st.session_state["session_id"],
                            "department_id": st.session_state["department_id"],
                            "template_id":   st.session_state["template_id"],
                            "answers":       answers,
                            "company":       st.session_state["company"]
                        }
                    )
                    if data:
                        st.session_state["document"]        = data
                        st.session_state["preview_content"] = None
                        st.session_state["show_regen"]      = False
                        st.session_state["val_data"]        = None
                        st.rerun()
                    else:
                        st.error(f"Error generating document: {err}")

        # Regenerate panel
        if st.session_state.get("show_regen") and doc:
            st.write("")
            with st.container(border=True):
                st.markdown("**What should be improved?**")
                feedback = st.text_area(
                    "", height=80, key="regen_text",
                    placeholder="Describe what to change — tone, content, structure…",
                    label_visibility="collapsed"
                )
                rc1, rc2 = st.columns(2)
                with rc1:
                    if st.button("✓ Confirm", type="primary",
                                  use_container_width=True, key="confirm_regen"):
                        with st.spinner("Regenerating…"):
                            data, err = api_post(
                                "/generate/regenerate",
                                {
                                    "session_id": st.session_state["session_id"],
                                    "answers":    st.session_state["answers"],
                                    "feedback":   feedback or None,
                                    "company":    st.session_state.get("company")
                                }
                            )
                            if data:
                                st.session_state["document"]   = data
                                st.session_state["show_regen"] = False
                                st.session_state["val_data"]   = None
                                st.rerun()
                            else:
                                st.error(f"Error: {err}")
                with rc2:
                    if st.button("✕ Cancel", use_container_width=True, key="cancel_regen"):
                        st.session_state["show_regen"] = False
                        st.rerun()

    # ════════════════════════════════════════
    # RIGHT COLUMN — document + actions
    # ════════════════════════════════════════
    with right_col:

        if not doc and not preview:
            st.write("")
            st.info("👈 Fill in the form and click **Generate Document** to see your document here.")

        else:
            # Document meta
            if preview:
                st.info("👁 Preview mode — not saved to library.")

            if doc and not preview:
                val = doc.get("validation_status", "pending")
                st.caption(
                    f"{status_badge(val)} {val.replace('_',' ').upper()}  ·  "
                    f"v{doc.get('version', '1.0')}  ·  ID {doc['document_id']}"
                )

            doc_title = (doc or {}).get("title", "")
            if doc_title:
                st.markdown(f"## {doc_title}")

            st.divider()

            # ── Actions row ──────────────────────────────────────────────────
            company   = st.session_state.get("company")
            dept_name = st.session_state.get("department_name", "")
            sections  = get_structured_sections(doc) if doc else []

            # Pre-compute bytes
            pdf_bytes  = generate_pdf_from_structured(sections, doc_title, company, dept_name) if sections else None
            docx_bytes = generate_docx_from_structured(sections, doc_title, company, dept_name) if sections else None

            act1, act2, act3, act4, act5, act6 = st.columns(6)

            with act1:
                if st.button("👁 Preview", use_container_width=True, key="prev_btn"):
                    with st.spinner("Generating preview…"):
                        pdata, perr = api_post(
                            "/generate/preview",
                            {
                                "department_id": st.session_state["department_id"],
                                "template_id":   st.session_state["template_id"],
                                "answers":       st.session_state["answers"],
                                "company":       company
                            }
                        )
                        if pdata:
                            st.session_state["preview_content"] = pdata.get("content", "")
                            st.rerun()
                        else:
                            st.error(f"Preview error: {perr}")

            with act2:
                if doc and st.button("🔁 Regenerate", use_container_width=True, key="regen_btn"):
                    st.session_state["show_regen"] = not st.session_state.get("show_regen", False)
                    st.rerun()

            with act3:
                # ── Validate — single click, simple result stored in session ─
                if doc and st.button("✅ Validate", use_container_width=True, key="val_btn"):
                    with st.spinner("Validating…"):
                        vdata, verr = api_post(
                            "/generate/validate",
                            {"document_id": doc["document_id"]}
                        )
                        if vdata:
                            st.session_state["val_data"] = vdata
                            new_status = "validated" if vdata.get("is_valid") else "needs_review"
                            st.session_state["document"]["validation_status"] = new_status
                            st.rerun()
                        else:
                            st.error(f"Validation error: {verr}")

            with act4:
                if pdf_bytes:
                    st.download_button(
                        "⬇ PDF", data=pdf_bytes,
                        file_name=f"{doc_title}.pdf",
                        mime="application/pdf",
                        use_container_width=True, key="dl_pdf"
                    )
                else:
                    st.button("⬇ PDF", disabled=True,
                              use_container_width=True, key="dl_pdf_dis")

            with act5:
                if docx_bytes:
                    st.download_button(
                        "⬇ DOCX", data=docx_bytes,
                        file_name=f"{doc_title}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        use_container_width=True, key="dl_docx"
                    )
                else:
                    st.button("⬇ DOCX", disabled=True,
                              use_container_width=True, key="dl_docx_dis")

            with act6:
                if doc and st.button("🚀 Notion", use_container_width=True, key="notion_btn"):
                    with st.spinner("Publishing…"):
                        ndata, nerr = api_post("/notion/publish", {"document_id": doc["document_id"]})
                        if ndata:
                            st.success(f"✅ Published! [Open in Notion]({ndata.get('notion_url', '')})")
                        else:
                            st.error(f"Notion error: {nerr}")

            # ── Validation result — clean, simple ────────────────────────────
            val_data = st.session_state.get("val_data")
            if val_data:
                st.write("")
                is_valid = val_data.get("is_valid", False)
                with st.container(border=True):
                    if is_valid:
                        st.success("✅ Document validated successfully.")
                    else:
                        score = val_data.get("llm_judge_result", {}).get("overall_score", "—")
                        rec   = val_data.get("llm_judge_result", {}).get("recommendation", "")
                        st.warning(f"⚠️ Document needs review — Quality score: {score}/10")
                        if rec:
                            st.caption(f"💡 {rec}")

            st.divider()

            # ── Document content ─────────────────────────────────────────────
            content = preview if preview else (doc.get("content", "") if doc else "")
            render_document(content)


# ═══════════════════════════════════════════
# PAGE: LIBRARY
# ═══════════════════════════════════════════

elif st.session_state["page"] == "library":

    # ── Document Detail View ─────────────────────────────────────────────────
    if st.session_state.get("library_doc"):
        lib_doc   = st.session_state["library_doc"]
        lib_title = lib_doc.get("title", "Document")
        company   = st.session_state.get("company")
        sections  = get_structured_sections(lib_doc)

        bc1, bc2 = st.columns([8, 2])
        with bc1:
            st.markdown(f"## {lib_title}")
            st.caption(
                f"v{lib_doc.get('version','1.0')}  ·  "
                f"{lib_doc.get('validation_status','').replace('_',' ').upper()}  ·  "
                f"{format_date(lib_doc.get('created_at',''))}"
            )
        with bc2:
            if st.button("← Back to Library", use_container_width=True, key="back_lib"):
                st.session_state["library_doc"] = None
                st.rerun()

        st.divider()
        render_document(lib_doc.get("content", ""))
        st.stop()

    # ── Library Main View ─────────────────────────────────────────────────────
    st.markdown("## 📚 Document Library")
    st.write("")

    # Filters
    f1, f2, f3 = st.columns(3)
    with f1:
        dept_data_lib    = api_get("/departments", show_error=False) or []
        dept_filter_map  = {d["name"]: d["id"] for d in dept_data_lib}
        filter_dept      = st.selectbox(
            "Department",
            ["All Departments"] + [d["name"] for d in dept_data_lib],
            key="lib_dept_f"
        )
    with f2:
        tmpl_lib = []
        if filter_dept != "All Departments":
            tmpl_lib = api_get(f"/templates/{dept_filter_map[filter_dept]}", show_error=False) or []
        else:
            for d in dept_data_lib:
                tmpl_lib.extend(api_get(f"/templates/{d['id']}", show_error=False) or [])
        tmpl_lib_display  = {short_name(t["name"]): t["id"] for t in tmpl_lib}
        filter_tmpl_disp  = st.selectbox(
            "Document Type",
            ["All Document Types"] + list(tmpl_lib_display.keys()),
            key="lib_tmpl_f"
        )
    with f3:
        filter_status = st.selectbox(
            "Status",
            ["All", "validated", "pending", "needs_review", "failed"],
            key="lib_status_f"
        )

    st.write("")

    params = {}
    if filter_dept != "All Departments":
        params["department_id"] = dept_filter_map[filter_dept]
    if filter_tmpl_disp != "All Document Types":
        params["template_id"] = tmpl_lib_display.get(filter_tmpl_disp)

    raw_docs = api_get("/documents", params=params, show_error=True)
    docs     = raw_docs or []

    if raw_docs is None:
        st.warning("⚠️ Could not load documents — check the error above.")
        st.stop()

    if filter_status != "All":
        docs = [d for d in docs if d.get("validation_status") == filter_status]

    if not docs:
        st.write("")
        st.info("No documents found. Generate your first document to see it here.")
    else:
        st.caption(f"{len(docs)} document(s) found")
        st.divider()

        company = st.session_state.get("company")

        for doc in docs:
            title   = doc.get("title", "Untitled")
            val     = doc.get("validation_status", "pending")
            version = doc.get("version", "1.0")
            created = format_date(doc.get("created_at", ""))
            doc_id  = doc["document_id"]
            sections = get_structured_sections(doc)

            # Pre-compute export bytes per row
            pdf_b  = generate_pdf_from_structured(sections, title, company, "") if sections else None
            docx_b = generate_docx_from_structured(sections, title, company, "") if sections else None

            # Row: title | view | pdf | docx | notion | delete
            c1, c2, c3, c4, c5, c6 = st.columns([4, 1, 1, 1, 1, 1])

            with c1:
                st.markdown(f"{status_badge(val)} **{title}**")
                st.caption(f"v{version}  ·  {created}  ·  {val.replace('_',' ').upper()}")

            with c2:
                if st.button("View", key=f"v_{doc_id}", use_container_width=True):
                    st.session_state["library_doc"] = doc
                    st.rerun()

            with c3:
                if pdf_b:
                    st.download_button(
                        "PDF", data=pdf_b,
                        file_name=f"{title}.pdf",
                        mime="application/pdf",
                        use_container_width=True, key=f"pdf_{doc_id}"
                    )

            with c4:
                if docx_b:
                    st.download_button(
                        "DOCX", data=docx_b,
                        file_name=f"{title}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        use_container_width=True, key=f"docx_{doc_id}"
                    )

            with c5:
                if st.button("Notion", key=f"notion_{doc_id}", use_container_width=True):
                    with st.spinner("Publishing…"):
                        ndata, nerr = api_post("/notion/publish", {"document_id": doc_id})
                        if ndata:
                            st.success(f"✅ [Open in Notion]({ndata.get('notion_url','')})")
                        else:
                            st.error(f"Error: {nerr}")

            with c6:
                if st.button("🗑", key=f"del_{doc_id}", use_container_width=True,
                              help="Delete this document"):
                    deleted = api_delete(f"/documents/{doc_id}")
                    if deleted:
                        st.success("Deleted.")
                        st.rerun()

            st.divider()

st.write("")
st.caption("DocForge Hub · Powered by AI")




